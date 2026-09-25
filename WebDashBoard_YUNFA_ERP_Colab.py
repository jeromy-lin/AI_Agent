# ===========================================================
#  主題 : WebBased AI 數據戰情室應用實作
#  目標 : 引導學員將 ERP 資料導入 Web 介面，運用 K-Means 
#         與多元線性迴歸，建立可查看分析結果與圖表的數據戰情室。
#  情境 : 以金屬機電／設備製造業為例，整合客戶、報價、
#         銷售訂單、產品及成本結算資料，進行客戶分群與
#         製造成本預測，作為業務經營及成本管理的參考。
#
#  作者 : 國立雲林科技大學電機工程系 林家仁
# ===========================================================


import os
import io
import sys
import time
import importlib
import subprocess
import warnings
from html import escape

warnings.filterwarnings("ignore")

# ------------------------------------------------------------
# Minimal dependency check
# No Transformers / Torch / GPU required.
# ------------------------------------------------------------
def _ensure_import(import_name, pip_spec=None):
    try:
        return importlib.import_module(import_name)
    except ImportError:
        if not pip_spec:
            return None
        subprocess.run(
            [sys.executable, "-m", "pip", "install", "-q", pip_spec],
            check=False
        )
        try:
            return importlib.import_module(import_name)
        except Exception:
            return None

for _module, _pip in [("numpy", "numpy"), ("pandas", "pandas"),
                       ("matplotlib", "matplotlib"), ("PIL", "pillow"),
                       ("sklearn", "scikit-learn"), ("openpyxl", "openpyxl"),
                       ("socksio", "socksio"),
                       ("gradio", "gradio>=5,<7")]:
    if _ensure_import(_module, _pip) is None:
        raise RuntimeError(f"必要套件 {_module} 無法安裝 / 載入。")

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, MaxNLocator
import gradio as gr

from PIL import Image
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.ensemble import IsolationForest
from sklearn.linear_model import LinearRegression
from sklearn.metrics import (
    r2_score,
    mean_absolute_error,
    mean_squared_error,
)
from sklearn.model_selection import GroupShuffleSplit
from sklearn.preprocessing import StandardScaler

print(f"✅ Gradio {gr.__version__}")
print("✅ CPU Analytics Environment Ready")
print("ℹ️ 本 Lab 不使用語言模型，不需要 GPU。")

def normalize_file_path(file_value):
    if file_value is None:
        return None
    if isinstance(file_value, (str, os.PathLike)):
        return os.fspath(file_value)
    if isinstance(file_value, dict):
        return file_value.get("path") or file_value.get("name")
    return getattr(file_value, "name", None)

def df_to_dark_html(df, empty_text="尚無資料"):
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return f"<div class='yunfa-table-empty'>{empty_text}</div>"

    safe = df.copy()
    numeric_cols = safe.select_dtypes(include="number").columns
    for col in numeric_cols:
        if col in ("訂單次數", "報價次數", "最近交易天數", "完工數量"):
            safe[col] = safe[col].map(lambda v: "—" if pd.isna(v) else f"{v:,.0f}")
        else:
            safe[col] = safe[col].map(lambda v: "—" if pd.isna(v) else f"{v:,.1f}")
    html = safe.to_html(
        index=False,
        escape=True,
        border=0,
        classes="yunfa-data-table",
    )
    return "<div class='yunfa-table-wrap'>" + html + "</div>"

def kpi_cards(kpis):
    html = "<div class='kpi-wrap'>"
    for key, value in kpis.items():
        html += (
            "<div class='kpi-card'>"
            f"<div class='kpi-label'>{escape(str(key))}</div>"
            f"<div class='kpi-value'>{escape(str(value))}</div>"
            "</div>"
        )
    html += "</div>"
    return html


def elbow_explanation_html(elbow=None):
    if elbow:
        values = dict(elbow)
        detail = f"目前示範設定為 <b>4 群</b>，其 SSE 為 <b>{values[4]:,.1f}</b>。"
        if 3 in values and 5 in values:
            detail += (f"由 3 群增至 4 群，SSE 降低 {values[3]-values[4]:,.1f}；"
                       f"由 4 群增至 5 群，SSE 降低 {values[4]-values[5]:,.1f}。")
    else:
        detail = "上傳 ERP 並執行分析後，可在圖表下拉選單查看 SSE 曲線。"
    return f"""
    <div class="metric-help">
      <div class="metric-help-title">📘 Elbow Method 怎麼看？</div>
      <div class="metric-help-row">
        對不同群數計算群內平方和（SSE）；群數增加時，SSE 通常會下降。
        觀察曲線由陡轉緩的位置，可作為選擇分群數的參考。
      </div>
      <div class="metric-help-row"><b>本次資料：</b>{detail}</div>
      <div class="metric-help-note">
        本程式固定以四群示範客戶策略；曲線不保證四群就是最佳群數，
        仍需結合各群人數及實際業務意義判讀。
      </div>
    </div>
    """


def regression_explanation_html(kpis):
    r2 = kpis.get("R²", "N/A")
    mae = kpis.get("MAE (KNTD)", "N/A")
    rmse = kpis.get("RMSE (KNTD)", "N/A")
    mape = kpis.get("MAPE(%)", "N/A")
    within3 = kpis.get("±3%內比例(%)", "N/A")

    return f"""
    <div class="metric-help">
      <div class="metric-help-title">📘 迴歸評估指標怎麼看？</div>

      <div class="metric-help-row">
        <span class="metric-help-name">R²</span>
        <span class="metric-help-current">{r2}</span><br>
        模型可解釋目標變異的程度。越接近 1 通常代表整體擬合越好。
      </div>

      <div class="metric-help-row">
        <span class="metric-help-name">MAE (KNTD)</span>
        <span class="metric-help-current">{mae}</span><br>
        平均絕對誤差。可直覺理解為：每張測試工單的預測成本，平均大約差多少 <b>KNTD（千元）</b>。
        <b>越低越好。</b>
      </div>

      <div class="metric-help-row">
        <span class="metric-help-name">RMSE (KNTD)</span>
        <span class="metric-help-current">{rmse}</span><br>
        均方根誤差，單位為 KNTD。會對特別大的預測誤差給更高懲罰。
        <b>越低越好。</b>
      </div>

      <div class="metric-help-row">
        <span class="metric-help-name">MAPE</span>
        <span class="metric-help-current">{mape}%</span><br>
        平均絕對百分比誤差。用百分比表示預測誤差，較容易跨不同成本規模理解。
        <b>越低越好。</b>
      </div>

      <div class="metric-help-row">
        <span class="metric-help-name">±3% 內比例</span>
        <span class="metric-help-current">{within3}%</span><br>
        測試資料中，預測成本與實際成本誤差落在 ±3% 以內的工單比例。
        <b>越高越好。</b>
      </div>

      <div class="metric-help-note">
        這些指標是用測試資料評估模型表現；每條產品線僅 18 筆，指標波動可能很大；R² 高不代表模型一定適合所有情境，
        還要一起查看 MAE、RMSE、MAPE、±3% 內比例與誤差分布圖。
      </div>
    </div>
    """

CUSTOM_CSS = """
:root { color-scheme: light; }
.gradio-container {
    max-width: 1440px !important;
    margin: auto !important;
    background: #fffdf9 !important;
    color: #273346 !important;
    font-family: "Noto Sans TC", "Microsoft JhengHei", sans-serif !important;
}
.yunfa-hero {
    padding: 24px 30px; margin: 4px 0 20px;
    border: 1px solid #e9ded1; border-radius: 18px;
    background: linear-gradient(115deg,#fff6e9,#f4f8f5);
}
.yunfa-hero h1 { margin: 0 0 8px; font-size: clamp(26px,3vw,36px); color: #263746; }
.yunfa-hero p { margin: 0; color: #63717d; font-size: 16px; }
.yunfa-section-title { font-size: 19px; font-weight: 800; margin: 16px 0 8px; color: #314555; }
.yunfa-unit-note { color: #64748b; margin: 5px 0 13px; font-size: 14px; }
.chart-meaning {
    margin: 12px 0 22px; padding: 17px 21px;
    border: 1px solid #dce7df; border-left: 5px solid #82b29b;
    border-radius: 12px; background: #f7fbf8; color: #334155;
    line-height: 1.75; font-size: 15px;
}
.chart-meaning strong { color: #2e5548; }
.chart-view img { object-fit: contain !important; }
.upload-instruction {
    border: 1px solid #e9ded1; border-radius: 14px;
    padding: 16px 20px; margin-bottom: 10px; background: #fff9f1;
}
.upload-instruction .file-kind {
    font-size: 19px; color: #594637;
    font-weight: 800;
    line-height: 1.35;
}
.upload-instruction .file-name {
    font-size: 17px;
    font-weight: 750;
    line-height: 1.4;
    margin-top: 4px;
}
.upload-instruction .file-note {
    font-size: 15px;
    color: #667484;
    margin-top: 4px;
}
.method-card {
    border: 1px solid #eadfd5; border-radius: 14px;
    padding: 18px 20px; min-height: 165px; background: #ffffff;
    box-shadow: 0 4px 16px rgba(83, 70, 56, .05);
}
.method-title {
    font-size: 19px; color: #36505b;
    font-weight: 800;
    margin-bottom: 6px;
}
.method-note {
    font-size: 15px;
    line-height: 1.7; color: #596978;
}
.yunfa-table-wrap {
    width: 100%;
    overflow-x: auto;
    border: 1px solid #e6e5e1;
    border-radius: 12px;
    background: #ffffff !important;
}
.yunfa-data-table {
    width: 100%;
    border-collapse: collapse;
    background: #ffffff !important;
    color: #334155 !important;
    font-size: 14px;
}
.yunfa-data-table thead th {
    background: #f5eee7 !important;
    color: #344454 !important;
    font-weight: 800 !important;
    border-right: 1px solid #eee8e1 !important;
    border-bottom: 1px solid #e3d8ca !important;
    padding: 10px 12px !important;
    white-space: nowrap;
    text-align: left;
}
.yunfa-data-table tbody td {
    background: #ffffff !important;
    color: #334155 !important;
    border-right: 1px solid #eeeef0 !important;
    border-bottom: 1px solid #eeeef0 !important;
    padding: 9px 12px !important;
    white-space: nowrap;
}
.yunfa-data-table tbody tr:nth-child(even) td {
    background: #fafbf9 !important;
}
.yunfa-data-table tbody tr:hover td {
    background: #fff3e4 !important;
    color: #263746 !important;
}
.yunfa-table-empty {
    border: 1px dashed #d9c9b9;
    border-radius: 10px;
    padding: 18px;
    background: #fffdf9 !important;
    color: #64748b !important;
    font-size: 15px;
}
.kpi-wrap {
    display: flex;
    gap: 12px;
    flex-wrap: wrap;
    margin: 8px 0 14px 0;
}
.kpi-card {
    border: 1px solid #eadfd5;
    border-radius: 12px;
    padding: 14px 18px;
    background: #fff9f2 !important;
    color: #0f172a !important;
    min-width: 155px; flex: 1 1 155px;
}
.kpi-card * {
    color: #0f172a !important;
}
.kpi-label {
    font-size: 13px;
    font-weight: 700;
}
.kpi-value {
    font-size: 22px;
    font-weight: 850;
    margin-top: 3px;
}

.metric-help {
    border: 1px solid #e2e8e2;
    border-radius: 12px;
    padding: 14px 16px;
    background: #f6faf7 !important;
    color: #0f172a !important;
    line-height: 1.55;
}
.metric-help * { color: #0f172a !important; }
.metric-help-title {
    font-size: 18px;
    font-weight: 850;
    margin-bottom: 8px;
}
.metric-help-row { margin: 7px 0; }
.metric-help-name {
    font-weight: 800;
    color: #0f3b78 !important;
}
.metric-help-current {
    display: inline-block;
    margin-left: 6px;
    padding: 2px 8px;
    border-radius: 999px;
    background: #e9f0eb !important;
    color: #0f172a !important;
    font-weight: 800;
}
.metric-help-note {
    margin-top: 10px;
    padding-top: 10px;
    border-top: 1px solid #dbe4ee;
    font-size: 13px;
    color: #475569 !important;
}
@media (max-width: 700px) {
    .yunfa-hero { padding: 18px; }
    .kpi-card { min-width: 130px; }
    .chart-view { max-height: 370px; }
}
"""

def _safe_num(df, cols):
    out = df.copy()
    for c in cols:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce")
    return out

def _fig_to_gallery(fig, caption):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=135, bbox_inches="tight", facecolor="white")
    buf.seek(0)
    img = Image.open(buf).convert("RGB").copy()
    plt.close(fig)
    return (img, caption)

CUSTOMER_CHART_CHOICES = [
    ("圖 1｜PCA 客戶分群位置", 0),
    ("圖 2｜各客戶群人數", 1),
    ("圖 3｜客戶群特徵比較", 2),
    ("圖 4｜Elbow Method 群數比較", 3),
]
COST_CHART_CHOICES = [
    ("圖 1｜實際與預測成本", 0),
    ("圖 2｜製造成本影響因素", 1),
    ("圖 3｜預測誤差分布", 2),
]

def chart_meaning(text):
    return f'<div class="chart-meaning"><strong>圖表意義與實務判讀</strong><br>{text}</div>'

def select_chart(charts, chart_index):
    if not charts:
        return None, chart_meaning("請先上傳 ERP 檔案並執行分析，再從上方選擇圖表。")
    index = int(chart_index) if chart_index is not None else 0
    if index < 0 or index >= len(charts):
        index = 0
    image, _caption, meaning = charts[index]
    return image, chart_meaning(meaning)

def _one_decimal_axis(ax, *, x=False, y=False):
    fmt = FuncFormatter(lambda value, _: f"{value:,.1f}")
    if x:
        ax.xaxis.set_major_formatter(fmt)
    if y:
        ax.yaxis.set_major_formatter(fmt)

# 原始匯出欄位 -> 教學 ERP 五張分析用工作表。附加來源欄位保留稽核線索。
RAW_REQUIRED = {
    "CRM_Account_Export": ["AccountNo", "AccountName", "CountryCode", "SegmentLevel", "CreatedOn", "IndustryGroup", "AccountStatus"],
    "QUOTE_LOG": ["QuoteNo", "CustCode", "MaterialCode", "QuoteCreated", "DiscPct", "EstMarginPct", "Outcome", "Currency"],
    "SD_SalesHistory": ["SalesDoc", "SoldTo", "PartNumber", "PostingDate", "OrderQty", "NetUnitPrice", "LineAmount", "Currency", "DocStatus"],
    "MM_ItemMaster": ["ItemNo", "ItemDescription", "BusinessFamily", "StdMaterialCost", "StdLaborCost", "StdSubcontractCost", "StdFactoryOH"],
    "CO_ActualCost": ["ManufacturingOrder", "Item", "CompletedQty", "CncHours", "AssemblyHours", "InspectionHours", "ChangeoverHours", "ActualMfgCost"],
}

def prepare_erp_tables(raw_path):
    """依原始值轉欄位，不估匯率、不推導不存在的客戶工單關聯。"""
    xls = pd.ExcelFile(raw_path)
    for sheet, columns in RAW_REQUIRED.items():
        if sheet not in xls.sheet_names:
            raise ValueError(f"原始資料缺少工作表：{sheet}")
        actual = pd.read_excel(xls, sheet)
        missing = set(columns) - set(actual.columns)
        if missing:
            raise ValueError(f"{sheet} 缺少欄位：{', '.join(sorted(missing))}")
    crm = pd.read_excel(xls, "CRM_Account_Export")
    quo = pd.read_excel(xls, "QUOTE_LOG")
    sales = pd.read_excel(xls, "SD_SalesHistory")
    item = pd.read_excel(xls, "MM_ItemMaster")
    cost = pd.read_excel(xls, "CO_ActualCost")
    if crm.AccountNo.duplicated().any() or item.ItemNo.duplicated().any():
        raise ValueError("客戶或產品主檔有重複主鍵，請先檢查來源。")
    if not quo.CustCode.isin(crm.AccountNo).all() or not sales.SoldTo.isin(crm.AccountNo).all():
        raise ValueError("報價／銷售有無法連結客戶主檔的 ID。")
    if not cost.Item.isin(item.ItemNo).all():
        raise ValueError("成本結算有無法連結產品主檔的 SKU。")
    if not np.allclose(pd.to_numeric(sales.OrderQty) * pd.to_numeric(sales.NetUnitPrice),
                       pd.to_numeric(sales.LineAmount)):
        raise ValueError("銷售明細的金額與數量 × 單價不符，請先核對來源。")
    tables = {
        "客戶主檔": crm.rename(columns={"AccountNo": "客戶ID", "AccountName": "客戶名稱",
            "CountryCode": "國別", "SegmentLevel": "客戶等級", "CreatedOn": "建立日期",
            "IndustryGroup": "產業別", "AccountStatus": "客戶狀態", "LegacyTag": "原系統標籤"}),
        "報價紀錄": quo.rename(columns={"QuoteNo": "報價ID", "CustCode": "客戶ID",
            "MaterialCode": "SKU", "QuoteCreated": "報價日期", "DiscPct": "折扣率",
            "EstMarginPct": "預估毛利率", "Outcome": "結果", "Currency": "幣別", "Channel": "通路"}),
        "銷售訂單": sales.rename(columns={"SalesDoc": "訂單ID", "SoldTo": "客戶ID",
            "PartNumber": "SKU", "PostingDate": "下單日期", "OrderQty": "訂購數量",
            "NetUnitPrice": "原幣成交單價", "LineAmount": "原幣訂單金額", "Currency": "幣別",
            "DocStatus": "單據狀態", "SalesRegion": "銷售區域"}),
        "產品主檔": item.rename(columns={"ItemNo": "SKU", "ItemDescription": "產品名稱",
            "BusinessFamily": "產品線", "StdMaterialCost": "標準材料成本_TWD",
            "StdLaborCost": "標準人工成本_TWD", "StdSubcontractCost": "標準委外成本_TWD",
            "StdFactoryOH": "標準製造費_TWD", "UOM": "單位", "ProcureType": "採購類型", "ABCClass": "ABC等級"}),
        "成本結算": cost.rename(columns={"ManufacturingOrder": "工單ID", "Item": "SKU",
            "CompletedQty": "完工數量", "CncHours": "CNC加工工時_hr",
            "AssemblyHours": "組裝工時_hr", "InspectionHours": "測試工時_hr",
            "ChangeoverHours": "換線工時_hr", "ActualMfgCost": "實際製造成本_TWD",
            "FiscalPeriod": "會計期間", "Plant": "廠區"}),
    }
    tables["報價紀錄"]["結果"] = tables["報價紀錄"]["結果"].replace({"Won": "成交", "Lost": "未成交", "Pending": "待確認"})
    tables["銷售訂單"]["成交單價_TWD"] = np.where(
        tables["銷售訂單"]["幣別"].eq("TWD"), tables["銷售訂單"]["原幣成交單價"], np.nan)
    tables["資料說明"] = pd.DataFrame([
        ["來源", "Test_Raw_Enterprise_Data.xlsx；五張分析表由同名英文來源表逐列轉換。"],
        ["客戶分析", "客戶主檔 + 報價紀錄 + 銷售訂單；報價率使用全部報價，採購金額僅使用 TWD 銷售。"],
        ["匯率", "來源無 USD→TWD 匯率；USD 單據以原幣保留，不納入 TWD 採購金額。"],
        ["成本分析", "成本結算依 SKU 連結產品主檔；來源未提供銷售訂單與製造工單的對應關係。"],
        ["模型", "四群 K-Means；各產品線線性回歸依工單ID切分訓練／測試，僅作教學示範。"],
    ], columns=["項目", "說明"])
    return tables

def write_erp_excel(tables, output_path):
    # Colab 執行時由本支完整程式產出上傳檔，無須額外程式碼儲存格。
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet, df in tables.items():
            df.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import PatternFill, Font
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="5B8F88")
                cell.font = Font(color="FFFFFF", bold=True)
    return output_path

def build_customer(erp_path):
    xls=pd.ExcelFile(erp_path)
    customer=pd.read_excel(xls,"客戶主檔"); quote=pd.read_excel(xls,"報價紀錄"); order=pd.read_excel(xls,"銷售訂單")
    quote["報價日期"]=pd.to_datetime(quote["報價日期"],errors="coerce"); order["下單日期"]=pd.to_datetime(order["下單日期"],errors="coerce")
    quote=_safe_num(quote,["折扣率","預估毛利率"]); order=_safe_num(order,["訂購數量","成交單價_TWD"])
    # 沒有外匯匯率，不能把 USD 單價誤當 TWD；保留 Excel 原始 USD 記錄。
    twd_order=order.loc[order["幣別"].eq("TWD")].copy()
    twd_order["訂單金額_TWD"]=twd_order["訂購數量"]*twd_order["成交單價_TWD"]
    ref=max(quote["報價日期"].max(),order["下單日期"].max())
    q=quote.groupby("客戶ID").agg(報價次數=("報價ID","count"),成交報價次數=("結果",lambda x:(x=="成交").sum()),平均折扣率=("折扣率","mean"),平均預估毛利率=("預估毛利率","mean"))
    q["報價成交率"]=q["成交報價次數"]/q["報價次數"].replace(0,np.nan)
    o=twd_order.groupby("客戶ID").agg(總採購金額_TWD=("訂單金額_TWD","sum"),訂單次數=("訂單ID","count"),平均訂單金額_TWD=("訂單金額_TWD","mean"),最近交易日=("下單日期","max"))
    period_years=max((ref-order["下單日期"].min()).days/365.25,1.0)
    o["最近交易天數"]=(ref-o["最近交易日"]).dt.days
    o["年化採購金額_TWD"]=o["總採購金額_TWD"]/period_years
    c=customer.merge(q.reset_index(),on="客戶ID",how="left").merge(o.reset_index(),on="客戶ID",how="left")
    if "客戶等級" in c.columns: c=c.rename(columns={"客戶等級":"既有客戶等級（ERP）"})
    feats=["年化採購金額_TWD","訂單次數","平均訂單金額_TWD","報價次數","報價成交率","平均折扣率","平均預估毛利率","最近交易天數"]
    c[feats]=c[feats].apply(pd.to_numeric,errors="coerce")
    c["最近交易天數"]=c["最近交易天數"].fillna((ref-order["下單日期"].min()).days+1)
    c[feats]=c[feats].fillna(0)
    if len(c)<5: raise ValueError("至少需要五位客戶才能建立四群及 PCA 圖表。")
    scaler=StandardScaler(); X=scaler.fit_transform(c[feats]); km=KMeans(n_clusters=4,random_state=42,n_init=20); c["Cluster"]=km.fit_predict(X)
    elbow=[(k, float(km.inertia_) if k==4 else float(KMeans(n_clusters=k,random_state=42,n_init=10).fit(X).inertia_))
           for k in range(1,min(8,len(c)-1)+1)]
    centers=pd.DataFrame(km.cluster_centers_,columns=feats,index=range(4)); rem=set(range(4)); names={}
    risk=(centers["最近交易天數"]-.5*centers["年化採購金額_TWD"]-.5*centers["訂單次數"]).idxmax(); names[risk]="沉睡／流失風險型"; rem.remove(risk)
    vip=(centers["年化採購金額_TWD"]+centers["訂單次數"]+centers["報價次數"]-centers["最近交易天數"]).loc[list(rem)].idxmax(); names[vip]="高價值活躍型"; rem.remove(vip)
    price=(centers["平均折扣率"]-centers["平均預估毛利率"]).loc[list(rem)].idxmax(); names[price]="價格敏感型"; rem.remove(price)
    names[list(rem)[0]]="成長潛力型"; c["K-Means行為分群"]=c["Cluster"].map(names)
    strategy={"高價值活躍型":"維持高互動、優先服務、交叉銷售與長期合作","成長潛力型":"提高成交率、增加產品組合與業務接觸","價格敏感型":"管理折扣、檢查毛利、採差異化報價策略","沉睡／流失風險型":"優先喚回、追蹤未成交原因與近期需求"}
    c["建議策略"]=c["K-Means行為分群"].map(strategy)
    pca=PCA(2); Z=pca.fit_transform(X); c["PC1"]=Z[:,0]; c["PC2"]=Z[:,1]
    return {"data":c,"features":feats,"elbow":elbow,"kpis":{"客戶數":len(c),"分群數":4}}

def build_cost(erp_path):
    xls=pd.ExcelFile(erp_path); product=pd.read_excel(xls,"產品主檔"); cost=pd.read_excel(xls,"成本結算")
    pk=["SKU","產品線","標準材料成本_TWD","標準人工成本_TWD","標準委外成本_TWD","標準製造費_TWD"]
    ck=["工單ID","SKU","完工數量","CNC加工工時_hr","組裝工時_hr","測試工時_hr","換線工時_hr","實際製造成本_TWD"]
    d=cost[ck].merge(product[pk],on="SKU",how="inner")
    for col in ck[2:]+pk[2:]:
        if col in d.columns: d[col]=pd.to_numeric(d[col],errors="coerce")
    d["標準材料總成本"]=d["完工數量"]*d["標準材料成本_TWD"]; d["標準人工總成本"]=d["完工數量"]*d["標準人工成本_TWD"]
    d["標準委外總成本"]=d["完工數量"]*d["標準委外成本_TWD"]; d["標準製造費總額"]=d["完工數量"]*d["標準製造費_TWD"]
    feats=["完工數量","標準材料總成本","標準人工總成本","標準委外總成本","標準製造費總額","CNC加工工時_hr","組裝工時_hr","測試工時_hr","換線工時_hr"]
    return {"data":d,"features":feats,"target":"實際製造成本_TWD","product_lines":sorted(d["產品線"].dropna().unique().tolist())}

def run_cost_model(cost_obj, product_line=None):
    d=cost_obj["data"]; feats=cost_obj["features"]; target=cost_obj["target"]
    if product_line not in set(d["產品線"].dropna()): product_line=cost_obj["product_lines"][0]
    s=d[d["產品線"]==product_line].dropna(subset=feats+[target]).copy()
    if len(s)<8: raise ValueError(f"{product_line} 建模資料不足。")
    X=s[feats]; y=s[target]
    splitter=GroupShuffleSplit(n_splits=1,test_size=.25,random_state=42)
    tr_idx,te_idx=next(splitter.split(X,y,groups=s["工單ID"]))
    Xtr,Xte,ytr,yte=X.iloc[tr_idx],X.iloc[te_idx],y.iloc[tr_idx],y.iloc[te_idx]
    scaler=StandardScaler(); Xtrz=scaler.fit_transform(Xtr); Xtez=scaler.transform(Xte); model=LinearRegression().fit(Xtrz,ytr); pred=model.predict(Xtez)
    r2=r2_score(yte,pred); mae=mean_absolute_error(yte,pred); rmse=np.sqrt(mean_squared_error(yte,pred)); mape=float(np.mean(np.abs((yte.values-pred)/np.maximum(np.abs(yte.values),1)))*100)
    res=s.loc[Xte.index,["工單ID","SKU","產品線"]].copy(); res["實際製造成本"]=yte.values; res["預測製造成本"]=pred; res["誤差率(%)"]=np.abs(res["實際製造成本"]-res["預測製造成本"])/np.maximum(res["實際製造成本"],1)*100
    res["是否超出±3%"] = np.where(res["誤差率(%)"]>3,"是","否")
    coef=pd.DataFrame({"變數":feats,"標準化迴歸係數":model.coef_}); coef["影響程度"]=coef["標準化迴歸係數"].abs(); coef=coef.sort_values("影響程度",ascending=False)
    return {"product_line":product_line,"result":res.sort_values("誤差率(%)",ascending=False),"coef":coef,"kpis":{"R²":float(r2),"MAE (KNTD)":float(mae)/1000,"RMSE (KNTD)":float(rmse)/1000,"MAPE(%)":mape,"±3%內比例(%)":float((res["誤差率(%)"]<=3).mean()*100)}}

def display_kpis(kpis):
    """Apply one-decimal formatting only at the UI boundary."""
    return {key: ("—" if value is None else
                  f"{value:,}" if key in ("客戶數", "分群數") else
                  f"{value:,.1f}") for key, value in kpis.items()}

def display_customer_table(data):
    table = data.copy()
    if "年化採購金額_TWD" in table:
        table["年化採購金額_TWD"] /= 1000
    rename = {"年化採購金額_TWD": "年化採購金額 (KNTD)"}
    for col in ("報價成交率", "平均折扣率", "平均預估毛利率"):
        if col in table:
            table[col] *= 100
            rename[col] = col + " (%)"
    return table.rename(columns=rename)

def display_cost_table(data):
    table = data.copy()
    for col in ("實際製造成本", "預測製造成本"):
        table[col] /= 1000
    return table.rename(columns={"實際製造成本": "實際製造成本 (KNTD)",
                                 "預測製造成本": "預測製造成本 (KNTD)",
                                 "誤差率(%)": "誤差率 (%)"})

def customer_chart_meanings(analysis):
    d = analysis["data"]
    group = "K-Means行為分群"
    counts = d[group].value_counts()
    n = len(d)
    annual = d.groupby(group)["年化採購金額_TWD"].sum()
    total_annual = annual.sum()
    vip = "高價值活躍型"
    dormant = "沉睡／流失風險型"
    price = "價格敏感型"
    vip_n = int(counts.get(vip, 0))
    dormant_n = int(counts.get(dormant, 0))
    dormant_value = float(annual.get(dormant, 0)) / 1000
    vip_share = float(annual.get(vip, 0)) / total_annual * 100 if total_annual else 0
    largest = str(counts.idxmax())
    largest_n = int(counts.max())
    discount = d.groupby(group)["平均折扣率"].mean()
    margin = d.groupby(group)["平均預估毛利率"].mean()
    price_discount = float(discount.get(price, 0)) * 100
    price_margin = float(margin.get(price, 0)) * 100
    return [
        f"本次將 {n:,} 位客戶分成四群。圖上相近的點代表交易與報價行為較相似；"
        f"高價值活躍型共 {vip_n:,} 位，年化採購金額占樣本 {vip_share:.1f}%。"
        "業務可先查看這一群的服務與續單狀況；若不同顏色大量重疊，分群不宜直接當作客戶分級依據。",

        f"人數最多的是「{escape(largest)}」（{largest_n:,} 位，占 {largest_n/n*100:.1f}%）；"
        f"沉睡／流失風險型 {dormant_n:,} 位，現有年化採購額約 {dormant_value:,.1f} KNTD。"
        "可用這個規模安排回訪名單，但不能僅憑人數推算未來流失金額或訂單機會。",

        f"各柱是群組相對全體平均的標準化差距，不是實際金額。這次價格敏感型的平均折扣率為 "
        f"{price_discount:.1f}%，平均預估毛利率為 {price_margin:.1f}%。"
        "報價檢討時可先核對該群折扣與毛利，再對照交易頻率及最近交易天數，決定哪些客戶要調整報價條件或優先追蹤。",

        "本圖比較不同分群數的群內平方和（SSE）。曲線若由明顯下降轉為平緩，轉折處可作為群數參考。"
        "圖中的四群為本課程的示範設定，不代表資料必然以四群最合適；可結合各群人數與業務可解釋性評估。",
    ]

def cost_chart_meanings(analysis):
    d = analysis["result"]
    n = len(d)
    error = d["誤差率(%)"]
    out = int((error > 3).sum())
    within = (n - out) / n * 100 if n else 0
    bias = float((d["預測製造成本"] - d["實際製造成本"]).mean()) / 1000
    bias_direction = "高估" if bias >= 0 else "低估"
    worst = d.loc[error.idxmax()]
    worst_id = escape(str(worst["工單ID"]))
    worst_error = float(worst["誤差率(%)"])
    median = float(error.median())
    p90 = float(error.quantile(.9))
    leading = analysis["coef"].iloc[0]
    variable = escape(str(leading["變數"]))
    coefficient = float(leading["標準化迴歸係數"]) / 1000
    exception_action = (
        "可先核對超標工單的材料、工時及委外結算；若偏差持續朝同一方向，報價前要另外檢查成本估計。"
        if out else "這批測試工單沒有超出 ±3% 的案例；可先將此結果當作目前樣本的基準，待後續工單持續驗證。"
    )
    error_action = (
        f"最大誤差為工單 {worst_id}（{worst_error:.1f}%）。可先查核該筆的製程、委外費與換線工時，"
        "再檢查高誤差工單是否集中於同一產品或作業條件。"
        if out else
        f"最大誤差為工單 {worst_id}（{worst_error:.1f}%），仍在 ±3% 以內；目前無需把它列為超標異常，"
        "後續可持續追蹤新工單的誤差分布。"
    )
    return [
        f"本次測試的 {n:,} 張工單中，{n-out:,} 張落在 ±3% 範圍內（{within:.1f}%），"
        f"{out:,} 張超出。平均預測成本較實際成本{bias_direction} {abs(bias):,.1f} KNTD。"
        f"{exception_action}",

        f"本次模型中相對影響最大的變數為「{variable}」，標準化係數 {coefficient:+,.1f} KNTD。"
        "這表示該變數每增加一個標準差時模型預測的變化量；請優先核對其成本與工時資料品質。"
        "成本項目彼此可能相關，此係數不能當作單獨增加該項費用的因果效果。",

        f"本次絕對誤差率中位數 {median:.1f}%，第 90 百分位 {p90:.1f}%；"
        f"超過 3% 的工單有 {out:,} 張。{error_action}",
    ]

def customer_original_charts(x):
    charts=[]
    meanings=customer_chart_meanings(x)
    d=x["data"]
    feats=x["features"]
    cluster_order=["高價值活躍型","成長潛力型","價格敏感型","沉睡／流失風險型"]
    name_en={
        "高價值活躍型":"High-Value Active",
        "成長潛力型":"Growth Potential",
        "價格敏感型":"Price Sensitive",
        "沉睡／流失風險型":"Dormant / Churn Risk",
    }
    colors={
        "高價值活躍型":"#F2C6A0",
        "成長潛力型":"#A8C69F",
        "價格敏感型":"#AFC8E6",
        "沉睡／流失風險型":"#D7C4E8",
    }

    # Original PCA
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    for group in cluster_order:
        part=d[d["K-Means行為分群"]==group]
        ax.scatter(part["PC1"],part["PC2"],s=85,alpha=.82,label=name_en[group],
                   color=colors[group],edgecolors="white",linewidths=.8)
    ax.set_xlabel("Principal Component 1"); ax.set_ylabel("Principal Component 2")
    ax.set_title("ERP Customer Segmentation - K-Means",fontsize=15,fontweight="bold")
    _one_decimal_axis(ax,x=True,y=True)
    ax.legend(frameon=False); ax.grid(True,alpha=.35); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"PCA 客戶分群位置"), meanings[0]))

    # Original cluster size
    size=d["K-Means行為分群"].value_counts().reindex(cluster_order).fillna(0)
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    bars=ax.bar([name_en[z] for z in size.index],size.values,
                color=[colors[z] for z in size.index],edgecolor="white",linewidth=1)
    for bar,val in zip(bars,size.values):
        ax.text(bar.get_x()+bar.get_width()/2,bar.get_height()+.25,f"{int(val)}",ha="center",va="bottom",fontsize=10)
    ax.set_xlabel("Customer Segment"); ax.set_ylabel("Number of Customers")
    ax.set_title("Customer Segment Distribution",fontsize=14,fontweight="bold")
    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
    ax.grid(axis="y",alpha=.35); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"各客戶群人數"), meanings[1]))

    # Original normalized profile
    scaler=StandardScaler()
    z=scaler.fit_transform(d[feats])
    zdf=pd.DataFrame(z,columns=feats,index=d.index)
    zdf["K-Means行為分群"]=d["K-Means行為分群"].values
    profile=zdf.groupby("K-Means行為分群")[feats].mean().reindex(cluster_order)
    profile_en={
        "年化採購金額_TWD":"Annual Purchase","訂單次數":"Order Frequency",
        "平均訂單金額_TWD":"Avg. Order","報價次數":"Quote Count",
        "報價成交率":"Conversion","平均折扣率":"Discount",
        "平均預估毛利率":"Gross Margin","最近交易天數":"Recency",
    }
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    xx=np.arange(len(feats)); width=.18
    for i,group in enumerate(cluster_order):
        ax.bar(xx+(i-1.5)*width,profile.loc[group].values,width=width,label=name_en[group],
               color=colors[group],edgecolor="white")
    ax.axhline(0,color="#888888",linewidth=1,linestyle="--")
    ax.set_xticks(xx); ax.set_xticklabels([profile_en[f] for f in feats],fontsize=9)
    ax.set_ylabel("Standardized Cluster Mean"); ax.set_title("Customer Segment Profile",fontsize=14,fontweight="bold")
    _one_decimal_axis(ax,y=True)
    ax.legend(frameon=False,ncol=2); ax.grid(axis="y",alpha=.35); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"客戶群特徵比較"), meanings[2]))

    k_values, sse_values = zip(*x["elbow"])
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    ax.plot(k_values,sse_values,color="#698F85",linewidth=2.5,marker="o",markersize=7)
    ax.scatter([4],[dict(x["elbow"])[4]],color="#D88B63",s=140,zorder=3,label="Teaching setting: K = 4")
    ax.set_xticks(k_values)
    ax.set_xlabel("Number of Clusters (K)")
    ax.set_ylabel("Within-Cluster Sum of Squares (SSE)")
    ax.set_title("Elbow Method for Customer Segmentation",fontsize=14,fontweight="bold")
    _one_decimal_axis(ax,y=True)
    ax.legend(frameon=False); ax.grid(True,alpha=.3); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"Elbow Method 群數比較"), meanings[3]))
    return charts

def cost_original_charts(r):
    charts=[]
    meanings=cost_chart_meanings(r)
    result=r["result"].copy()
    coef=r["coef"].copy()
    product_line=r["product_line"]
    product_line_en={
        "汽車精密零組件":"Automotive Precision Parts",
        "精密金屬加工件":"Precision Metal Components",
        "機電／自動化模組":"Mechatronics & Automation Modules",
        "半導體設備零組件":"Semiconductor Equipment Parts",
    }.get(product_line,str(product_line))
    feature_en={
        "完工數量":"Completed Quantity",
        "標準材料總成本":"Standard Material Cost",
        "標準人工總成本":"Standard Labor Cost",
        "標準委外總成本":"Standard Outsourcing Cost",
        "標準製造費總額":"Standard Manufacturing Overhead",
        "CNC加工工時_hr":"CNC Machining Hours",
        "組裝工時_hr":"Assembly Hours",
        "測試工時_hr":"Test Hours",
        "換線工時_hr":"Setup Hours",
    }

    # Original CHART 1 Actual vs Predicted with ±3%
    actual=result["實際製造成本"]/1000
    pred=result["預測製造成本"]/1000
    within=result["誤差率(%)"]<=3
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    ax.scatter(actual[within],pred[within],alpha=.82,s=62,color="#F6B26B",edgecolors="white",linewidths=.7,label="Within ±3%")
    ax.scatter(actual[~within],pred[~within],alpha=.95,s=78,color="#D9534F",edgecolors="white",linewidths=.9,label="Outside ±3%")
    min_v=min(actual.min(),pred.min()); max_v=max(actual.max(),pred.max())
    ax.plot([min_v,max_v],[min_v,max_v],linestyle="--",linewidth=2,color="#6FA8DC",label="Perfect Prediction")
    xb=np.linspace(min_v,max_v,200); lo=xb*.97; hi=xb*1.03
    ax.fill_between(xb,lo,hi,alpha=.18,color="#93C47D",label="±3% Prediction Range")
    ax.plot(xb,lo,linestyle=":",linewidth=1.2,color="#93C47D")
    ax.plot(xb,hi,linestyle=":",linewidth=1.2,color="#93C47D")
    ax.set_xlabel("Actual Manufacturing Cost (KNTD)")
    ax.set_ylabel("Predicted Manufacturing Cost (KNTD)")
    ax.set_title(f"Actual vs Predicted Cost with ±3% Range - {product_line_en}",fontweight="bold")
    _one_decimal_axis(ax,x=True,y=True)
    ax.grid(True,alpha=.35); ax.legend(frameon=False); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"實際與預測成本"), meanings[0]))

    # Original CHART 2 Cost Drivers
    pc=coef.sort_values("影響程度",ascending=True).copy()
    pc["English"]=pc["變數"].map(feature_en).fillna(pc["變數"])
    pc["CoefK"]=pc["標準化迴歸係數"]/1000
    palette=["#A8C69F","#F2C6A0","#AFC8E6","#D7C4E8","#E8D58A","#C9D8B6","#E6B8AF","#B7D7D0","#F4D7B9"]
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    bars=ax.barh(pc["English"],pc["CoefK"],color=[palette[i%len(palette)] for i in range(len(pc))],
                 edgecolor="#F7F2EC",linewidth=1.2,height=.62)
    coefmax=max(pc["CoefK"].abs().max(),1); off=.02*coefmax
    for bar,val in zip(bars,pc["CoefK"]):
        ax.text(val+(off if val>=0 else -off),bar.get_y()+bar.get_height()/2,f"{val:,.1f}",
                va="center",ha="left" if val>=0 else "right",fontsize=9,color="#3F3F3F")
    ax.axvline(0,linestyle="--",linewidth=1.5,color="#888888")
    xmin=min(pc["CoefK"].min(),0); xmax=max(pc["CoefK"].max(),0); pad=.12*max(abs(xmin),abs(xmax),1)
    ax.set_xlim(xmin-pad,xmax+pad)
    ax.set_xlabel("Standardized Coefficient (KNTD)"); ax.set_ylabel("Feature")
    ax.set_title(f"Manufacturing Cost Drivers - {product_line_en}",fontweight="bold")
    _one_decimal_axis(ax,x=True)
    ax.grid(axis="x",alpha=.35); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"製造成本影響因素"), meanings[1]))

    # Original CHART 3 Error Distribution
    fig,ax=plt.subplots(figsize=(13,6.1),facecolor="white")
    ax.hist(result["誤差率(%)"],bins=15,color="#B7D7C4",edgecolor="white",linewidth=1)
    ax.set_xlabel("Absolute Error Rate (%)"); ax.set_ylabel("Number of Work Orders")
    ax.set_title(f"Prediction Error Distribution - {product_line_en}",fontweight="bold")
    _one_decimal_axis(ax,x=True)
    ax.yaxis.set_major_locator(MaxNLocator(integer=True))
    ax.grid(axis="y",alpha=.35); fig.tight_layout()
    charts.append((*_fig_to_gallery(fig,"預測誤差分布"), meanings[2]))
    return charts


# ------------------------------------------------------------
# ERP Manual Analytics UI
# ------------------------------------------------------------
def get_product_lines(erp_file):
    path = normalize_file_path(erp_file)
    if not path:
        return gr.update(choices=[], value=None), "請上傳 YUNFA_ERP.xlsx，分析會自動開始。"
    try:
        xls = pd.ExcelFile(path)
        required = {"客戶主檔", "報價紀錄", "銷售訂單", "產品主檔", "成本結算"}
        missing = required - set(xls.sheet_names)
        if missing:
            return (gr.update(choices=[], value=None),
                    "❌ ERP 缺少工作表：" + "、".join(sorted(missing)))
        product = pd.read_excel(xls, "產品主檔")
        lines = sorted(product["產品線"].dropna().astype(str).unique().tolist())
        return gr.update(
            choices=lines,
            value=lines[0] if lines else None,
        ), f"✅ 已讀取 ERP：{len(xls.sheet_names)} 張工作表、{len(lines)} 條產品線。正在執行兩項分析，結果會出現在下方分頁。"
    except Exception as exc:
        return (gr.update(choices=[], value=None),
                f"❌ ERP 讀取失敗：{type(exc).__name__}: {exc}")

def run_customer_web(erp_file, progress=gr.Progress()):
    path = normalize_file_path(erp_file)
    if not path:
        return (
            "❌ 請先上傳 YUNFA_ERP.xlsx。",
            "",
            elbow_explanation_html(None),
            df_to_dark_html(pd.DataFrame(), "尚無客戶分群結果"),
            [], gr.update(value=0), None,
            chart_meaning("請先上傳 ERP 檔案並執行分析。")
        )

    t0 = time.time()
    try:
        progress(0.18, desc="讀取客戶 / 報價 / 訂單")
        result = build_customer(path)
        progress(0.65, desc="執行 K-Means Customer Segmentation")

        table = result["data"].copy()
        cols = [
            c for c in [
                "客戶ID","客戶名稱","既有客戶等級（ERP）",
                "年化採購金額_TWD","訂單次數","報價成交率",
                "平均折扣率","平均預估毛利率","最近交易天數",
                "K-Means行為分群","建議策略"
            ] if c in table.columns
        ]
        table = table[cols].sort_values("年化採購金額_TWD", ascending=False).head(20)
        table = display_customer_table(table)

        progress(0.83, desc="產生 4 張客戶分群圖表")
        charts = customer_original_charts(result)

        elapsed = time.time() - t0
        status = (
            "✅ K-Means Customer Segmentation 完成  \n"
            f"CPU 執行時間：**{elapsed:.1f} 秒**  \n"
            "K-Means 使用交易 / 報價行為做分群；ERP 原始客戶等級不作為分群輸入。"
        )
        progress(1.0, desc="完成")
        return (
            status,
            kpi_cards(display_kpis(result["kpis"])),
            elbow_explanation_html(result["elbow"]),
            df_to_dark_html(table),
            charts, gr.update(value=0), *select_chart(charts, 0)
        )
    except Exception as e:
        return (
            f"❌ K-Means 執行失敗：{type(e).__name__}: {e}",
            "",
            elbow_explanation_html(None),
            df_to_dark_html(pd.DataFrame(), "K-Means 執行失敗"),
            [], gr.update(value=0), None,
            chart_meaning("客戶分群未完成，請檢查上方錯誤訊息。")
        )

def run_cost_web(erp_file, product_line, progress=gr.Progress()):
    path = normalize_file_path(erp_file)
    if not path:
        return (
            "❌ 請先上傳 YUNFA_ERP.xlsx。",
            "",
            regression_explanation_html({}),
            df_to_dark_html(pd.DataFrame(), "尚無成本模型結果"),
            [], gr.update(value=0), None,
            chart_meaning("請先上傳 ERP 檔案並執行分析。")
        )

    t0 = time.time()
    try:
        progress(0.18, desc="讀取產品與成本結算資料")
        base = build_cost(path)
        if not product_line:
            product_line = base["product_lines"][0] if base["product_lines"] else None

        progress(0.60, desc=f"建立 {product_line} Multiple Regression")
        result = run_cost_model(base, product_line)

        table = display_cost_table(result["result"].copy().head(15))

        progress(0.83, desc="產生 3 張成本分析圖表")
        charts = cost_original_charts(result)

        elapsed = time.time() - t0
        status = (
            f"✅ Multiple Linear Regression 完成｜{result['product_line']}  \n"
            f"CPU 執行時間：**{elapsed:.1f} 秒**  \n"
            "同一工單不跨訓練／測試；每條產品線僅 18 筆，結果僅供教學。標準化係數不代表因果關係。"
        )
        progress(1.0, desc="完成")
        return (
            status,
            kpi_cards(display_kpis(result["kpis"])),
            regression_explanation_html(display_kpis(result["kpis"])),
            df_to_dark_html(table),
            charts, gr.update(value=0), *select_chart(charts, 0)
        )
    except Exception as e:
        return (
            f"❌ Multiple Regression 執行失敗：{type(e).__name__}: {e}",
            "",
            regression_explanation_html({}),
            df_to_dark_html(pd.DataFrame(), "Regression 執行失敗"),
            [], gr.update(value=0), None,
            chart_meaning("製造成本分析未完成，請檢查上方錯誤訊息。")
        )

def create_erp_app():
    with gr.Blocks(
        title="WebBased AI 數據戰情室應用實作",
        css=CUSTOM_CSS,
    ) as demo:
        gr.HTML("""<div class="yunfa-hero"><h1>WebBased AI 數據戰情室應用實作</h1>
                <p>YUNFA ERP｜客戶分群 · 製造成本預測｜K-Means + Multiple Linear Regression</p></div>""")

        gr.HTML(
            """
            <div class="upload-instruction">
              <div class="file-kind">📘 ERP 商務 / 成本資料</div>
              <div class="file-name">請上傳：YUNFA_ERP_本次整理.xlsx</div>
              <div class="file-note">上傳後自動執行兩項分析，結果顯示於下方分頁。金額顯示 KNTD（千元）、小數一位；可再按執行鈕重新分析。</div>
            </div>
            """
        )
        erp_file = gr.File(
            label="選擇已整理的 ERP Excel（可更換檔案）",
            type="filepath",
            file_types=[".xlsx"],
            value=PREPARED_ERP_PATH,
        )
        upload_status = gr.Markdown("請上傳 YUNFA_ERP.xlsx，分析會自動開始。")

        with gr.Row():
            with gr.Column():
                gr.HTML(
                    """
                    <div class="method-card">
                      <div class="method-title">方法 1｜K-Means Customer Segmentation</div>
                      <div class="method-note">
                        問題：哪些客戶行為相似？應採取不同經營策略嗎？<br>
                        使用：採購金額、訂單頻率、成交率、折扣、毛利、Recency 等 8 個特徵。<br>
                        輸出：PCA、Segment Distribution、Segment Profile、Elbow Method，共 4 張圖。
                      </div>
                    </div>
                    """
                )
                run_kmeans = gr.Button("▶ 執行 K-Means 客戶分群", variant="primary")

            with gr.Column():
                gr.HTML(
                    """
                    <div class="method-card">
                      <div class="method-title">方法 2｜Multiple Linear Regression</div>
                      <div class="method-note">
                        問題：製造成本可以由哪些成本與工時因素解釋 / 預測？<br>
                        使用：標準成本、完工數量、CNC / 組裝 / 測試 / 換線工時。<br>
                        輸出：Actual vs Predicted、Cost Drivers、Error Distribution，共 3 張圖。
                      </div>
                    </div>
                    """
                )
                product_line = gr.Dropdown(
                    choices=[],
                    label="選擇產品線",
                    info="上傳 ERP 後會自動讀取產品線",
                )
                run_reg = gr.Button("▶ 執行 Multiple Regression", variant="secondary")

        with gr.Tab("K-Means 客戶分群結果"):
            customer_status = gr.Markdown()
            with gr.Row():
                with gr.Column(scale=2):
                    customer_kpis = gr.HTML()
                with gr.Column(scale=1):
                    customer_help = gr.HTML(value=elbow_explanation_html(None))
            gr.HTML('<div class="yunfa-section-title">客戶分群明細</div><div class="yunfa-unit-note">金額：KNTD（千元）｜比率：%｜顯示至小數一位</div>')
            customer_table = gr.HTML(value=df_to_dark_html(pd.DataFrame(), "尚未執行 K-Means"))
            gr.HTML('<div class="yunfa-section-title">分群圖表</div>')
            customer_chart_choice = gr.Dropdown(
                choices=CUSTOMER_CHART_CHOICES, value=0, label="選擇要查看的分析圖表",
                info="一次顯示一張圖及其意義", interactive=True,
            )
            customer_charts = gr.State([])
            customer_chart_image = gr.Image(
                label="客戶分群圖表", type="pil", interactive=False,
                height=540, elem_classes=["chart-view"],
            )
            customer_chart_meaning = gr.HTML(value=chart_meaning("請先執行客戶分群分析。"))

        with gr.Tab("Multiple Regression 結果"):
            cost_status = gr.Markdown()
            with gr.Row():
                with gr.Column(scale=2):
                    cost_kpis = gr.HTML()
                with gr.Column(scale=1):
                    cost_help = gr.HTML(value=regression_explanation_html({}))
            gr.HTML('<div class="yunfa-section-title">工單預測明細</div><div class="yunfa-unit-note">製造成本：KNTD（千元）｜誤差率：%｜顯示至小數一位</div>')
            cost_table = gr.HTML(value=df_to_dark_html(pd.DataFrame(), "尚未執行 Regression"))
            gr.HTML('<div class="yunfa-section-title">成本分析圖表</div>')
            cost_chart_choice = gr.Dropdown(
                choices=COST_CHART_CHOICES, value=0, label="選擇要查看的分析圖表",
                info="一次顯示一張圖及其意義", interactive=True,
            )
            cost_charts = gr.State([])
            cost_chart_image = gr.Image(
                label="製造成本圖表", type="pil", interactive=False,
                height=540, elem_classes=["chart-view"],
            )
            cost_chart_meaning = gr.HTML(value=chart_meaning("請先執行製造成本分析。"))

        run_kmeans.click(
            run_customer_web,
            inputs=[erp_file],
            outputs=[customer_status, customer_kpis, customer_help, customer_table,
                     customer_charts, customer_chart_choice, customer_chart_image, customer_chart_meaning],
        )
        run_reg.click(
            run_cost_web,
            inputs=[erp_file, product_line],
            outputs=[cost_status, cost_kpis, cost_help, cost_table,
                     cost_charts, cost_chart_choice, cost_chart_image, cost_chart_meaning],
        )
        customer_chart_choice.change(
            select_chart, inputs=[customer_charts, customer_chart_choice],
            outputs=[customer_chart_image, customer_chart_meaning],
        )
        cost_chart_choice.change(
            select_chart, inputs=[cost_charts, cost_chart_choice],
            outputs=[cost_chart_image, cost_chart_meaning],
        )
        erp_file.change(
            get_product_lines, inputs=[erp_file], outputs=[product_line, upload_status],
        ).then(
            run_customer_web, inputs=[erp_file],
            outputs=[customer_status, customer_kpis, customer_help, customer_table,
                     customer_charts, customer_chart_choice, customer_chart_image, customer_chart_meaning],
        ).then(
            run_cost_web, inputs=[erp_file, product_line],
            outputs=[cost_status, cost_kpis, cost_help, cost_table,
                     cost_charts, cost_chart_choice, cost_chart_image, cost_chart_meaning],
        )
        return demo

PREPARED_ERP_PATH = None

def setup_colab_data():
    """單儲存格入口：上傳原始 Excel → 整理 → 提供下載 → 啟動戰情室。"""
    try:
        from google.colab import files
    except ImportError as exc:
        raise RuntimeError("請在 Google Colab 中執行本完整程式碼並上傳原始 Excel。") from exc
    print("📤 請選擇 Test_Raw_Enterprise_Data.xlsx（本次原始資料）")
    uploaded = files.upload()
    candidates = [name for name in uploaded if name.lower().endswith(".xlsx")]
    if len(candidates) != 1:
        raise ValueError("請一次只上傳一份原始 .xlsx，然後重新執行此儲存格。")
    tables = prepare_erp_tables(candidates[0])
    output_path = os.path.abspath("YUNFA_ERP_本次整理.xlsx")
    write_erp_excel(tables, output_path)
    print("✅ 整理完成：" + ", ".join(f"{name} {len(df)} 筆" for name, df in tables.items() if name != "資料說明"))
    customers = build_customer(output_path)
    costs = build_cost(output_path)
    for line in costs["product_lines"]:
        run_cost_model(costs, line)
    print(f"✅ 分群 {len(customers['data'])} 位客戶；成本回歸 {len(costs['data'])} 筆、{len(costs['product_lines'])} 條產品線。")
    print("📥 正在提供整理後 Excel 下載；關閉下載提示後可使用下方戰情室。")
    files.download(output_path)
    return output_path

if __name__ == "__main__":
    PREPARED_ERP_PATH = setup_colab_data()
    demo = create_erp_app()
    print("🚀 啟動 YUNFA ERP Analytics Lab")
    demo.queue().launch(share=True, show_error=True, debug=False,
                        allowed_paths=[PREPARED_ERP_PATH])
